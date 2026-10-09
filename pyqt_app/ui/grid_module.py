# -*- coding: utf-8 -*-
"""页签1：网格生成（点云 → 施工网格数据表）"""
import os

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QFrame, QLineEdit, QDoubleSpinBox, QPlainTextEdit, QFileDialog,
    QMessageBox, QGridLayout,
)
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QPixmap

from app_config import (
    GRID_DIR, SCRIPT_GRID_DEPTH, algo_python, save_config,
)
from utils.subprocess_runner import CommandRunner


class ScaledImageLabel(QLabel):
    """保持宽高比缩放显示图片，窗口 resize 时自动重缩放"""

    def __init__(self, placeholder="暂无图片", parent=None):
        super().__init__(parent)
        self._pixmap = None
        self._placeholder = placeholder
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setMinimumSize(240, 180)
        self.setStyleSheet("background:#0A0F1A; border:1px solid #334155; "
                           "border-radius:8px; color:#64748B;")
        self.setText(placeholder)

    def set_image(self, path):
        pix = QPixmap(path)
        if pix.isNull():
            self._pixmap = None
            self.setText(self._placeholder)
            return
        self._pixmap = pix
        self._rescale()

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._rescale()

    def _rescale(self):
        if self._pixmap is None or self.width() < 10:
            return
        scaled = self._pixmap.scaled(
            self.size(), Qt.AspectRatioMode.KeepAspectRatio,
            Qt.TransformationMode.SmoothTransformation)
        self.setPixmap(scaled)


class StatCard(QFrame):
    """统计卡片"""

    def __init__(self, title, parent=None):
        super().__init__(parent)
        self.setObjectName("card")
        layout = QVBoxLayout(self)
        layout.setContentsMargins(14, 10, 14, 10)
        layout.setSpacing(2)
        t = QLabel(title)
        t.setStyleSheet("color:#94A3B8; font-size:12px; border:none;")
        layout.addWidget(t)
        self.value = QLabel("--")
        self.value.setStyleSheet("color:#F1F5F9; font-size:20px; font-weight:bold; border:none;")
        layout.addWidget(self.value)

    def set_value(self, text):
        self.value.setText(str(text))


class GridModule(QWidget):
    """页签1：网格生成"""

    def __init__(self, main_window):
        super().__init__()
        self.main = main_window
        self.runner = None
        self._build_ui()
        self._load_existing_outputs()

    # ==================== UI ====================

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setContentsMargins(20, 16, 20, 16)
        layout.setSpacing(12)

        # ---- 参数卡片 ----
        card = QFrame()
        card.setObjectName("card")
        form = QGridLayout(card)
        form.setContentsMargins(16, 12, 16, 12)
        form.setHorizontalSpacing(10)
        form.setVerticalSpacing(8)

        gcfg = self.main.config.get("grid", {})

        self.precise_edit = QLineEdit(gcfg.get("precise", ""))
        form.addWidget(QLabel("精确区点云:"), 0, 0)
        form.addWidget(self.precise_edit, 0, 1, 1, 3)
        btn = QPushButton("浏览…")
        btn.clicked.connect(lambda: self._browse_file(self.precise_edit, "选择精确区点云文件"))
        form.addWidget(btn, 0, 4)

        self.full_edit = QLineEdit(gcfg.get("full", ""))
        form.addWidget(QLabel("完整区域点云:"), 1, 0)
        form.addWidget(self.full_edit, 1, 1, 1, 3)
        btn = QPushButton("浏览…")
        btn.clicked.connect(lambda: self._browse_file(self.full_edit, "选择完整区域点云文件"))
        form.addWidget(btn, 1, 4)

        self.cell_x_spin = QDoubleSpinBox()
        self.cell_x_spin.setRange(0.1, 100.0)
        self.cell_x_spin.setDecimals(2)
        self.cell_x_spin.setSingleStep(0.1)
        self.cell_x_spin.setValue(float(gcfg.get("cell_x", "4.2") or 4.2))
        self.cell_y_spin = QDoubleSpinBox()
        self.cell_y_spin.setRange(0.1, 100.0)
        self.cell_y_spin.setDecimals(2)
        self.cell_y_spin.setSingleStep(0.1)
        self.cell_y_spin.setValue(float(gcfg.get("cell_y", "7.5") or 7.5))
        form.addWidget(QLabel("网格宽度 X (m):"), 2, 0)
        form.addWidget(self.cell_x_spin, 2, 1)
        form.addWidget(QLabel("网格长度 Y (m):"), 2, 2)
        form.addWidget(self.cell_y_spin, 2, 3)

        self.out_dir_edit = QLineEdit(gcfg.get("out_dir", GRID_DIR))
        form.addWidget(QLabel("输出目录:"), 3, 0)
        form.addWidget(self.out_dir_edit, 3, 1, 1, 3)
        btn = QPushButton("浏览…")
        btn.clicked.connect(self._browse_dir)
        form.addWidget(btn, 3, 4)

        layout.addWidget(card)

        # ---- 按钮行 ----
        btn_row = QHBoxLayout()
        self.run_btn = QPushButton("▶ 开始生成")
        self.run_btn.setObjectName("primaryBtn")
        self.run_btn.clicked.connect(self._run)
        btn_row.addWidget(self.run_btn)

        self.open_btn = QPushButton("📂 打开输出目录")
        self.open_btn.clicked.connect(self._open_out_dir)
        btn_row.addWidget(self.open_btn)
        btn_row.addStretch()
        layout.addLayout(btn_row)

        # ---- 统计卡片 ----
        stats_row = QHBoxLayout()
        stats_row.setSpacing(10)
        self.stat_total = StatCard("总网格数")
        self.stat_precise = StatCard("精确区网格")
        self.stat_filled = StatCard("估计区网格")
        self.stat_volume = StatCard("总填方量 (m³)")
        for c in (self.stat_total, self.stat_precise, self.stat_filled, self.stat_volume):
            stats_row.addWidget(c, 1)
        layout.addLayout(stats_row)

        # ---- 图片 + 日志 ----
        bottom = QHBoxLayout()
        bottom.setSpacing(12)

        img_col = QVBoxLayout()
        self.zone_img = ScaledImageLabel("区域/深度图 (zone_and_depth.png)")
        self.depth_img = ScaledImageLabel("深度网格图 (depth_grid.png)")
        img_col.addWidget(self.zone_img, 1)
        img_col.addWidget(self.depth_img, 1)
        bottom.addLayout(img_col, 3)

        log_col = QVBoxLayout()
        log_col.addWidget(QLabel("运行日志"))
        self.log = QPlainTextEdit()
        self.log.setReadOnly(True)
        log_col.addWidget(self.log, 1)
        bottom.addLayout(log_col, 2)

        layout.addLayout(bottom, 1)

    # ==================== 逻辑 ====================

    def _browse_file(self, edit, title):
        path, _ = QFileDialog.getOpenFileName(
            self, title, GRID_DIR, "文本文件 (*.txt);;所有文件 (*)")
        if path:
            edit.setText(path)

    def _browse_dir(self):
        path = QFileDialog.getExistingDirectory(self, "选择输出目录", GRID_DIR)
        if path:
            self.out_dir_edit.setText(path)

    def _open_out_dir(self):
        out_dir = self.out_dir_edit.text().strip() or GRID_DIR
        os.makedirs(out_dir, exist_ok=True)
        if hasattr(os, "startfile"):
            os.startfile(out_dir)

    def _run(self):
        if self.runner is not None and self.runner.isRunning():
            return
        precise = self.precise_edit.text().strip()
        full = self.full_edit.text().strip()
        out_dir = self.out_dir_edit.text().strip()
        for path, name in ((precise, "精确区点云"), (full, "完整区域点云")):
            if not os.path.isfile(path):
                QMessageBox.warning(self, "参数错误", "%s文件不存在:\n%s" % (name, path))
                return
        if not out_dir:
            QMessageBox.warning(self, "参数错误", "请填写输出目录")
            return

        gcfg = self.main.config.setdefault("grid", {})
        gcfg.update({
            "precise": precise, "full": full,
            "cell_x": str(self.cell_x_spin.value()),
            "cell_y": str(self.cell_y_spin.value()),
            "out_dir": out_dir,
        })
        save_config(self.main.config)

        cmd = [algo_python(), SCRIPT_GRID_DEPTH,
               "--precise", precise, "--full", full,
               "--cell-x", str(self.cell_x_spin.value()),
               "--cell-y", str(self.cell_y_spin.value()),
               "--out-dir", out_dir]

        self.log.clear()
        self.run_btn.setEnabled(False)
        self.run_btn.setText("⏳ 计算中...")
        self.main.set_status("网格生成中…")

        self.runner = CommandRunner(cmd, cwd=GRID_DIR, parent=self)
        self.runner.log_line.connect(self._append_log)
        self.runner.finished_rc.connect(self._on_finished)
        self.runner.start()

    def _append_log(self, line):
        self.log.appendPlainText(line)

    def _on_finished(self, rc):
        self.run_btn.setEnabled(True)
        self.run_btn.setText("▶ 开始生成")
        if rc == 0:
            self.main.set_status("网格生成完成")
            self._load_existing_outputs()
        else:
            self.main.set_status("网格生成失败 (exit %d)" % rc, ok=False)
            self._append_log("!!! 进程退出码 %d" % rc)

    def _load_existing_outputs(self):
        """加载已有的统计与图片产物（启动时/运行完成后调用）"""
        out_dir = self.out_dir_edit.text().strip() or GRID_DIR
        self._load_stats(os.path.join(out_dir, "施工网格数据表.xlsx"))
        self.zone_img.set_image(os.path.join(out_dir, "zone_and_depth.png"))
        self.depth_img.set_image(os.path.join(out_dir, "depth_grid.png"))

    def _load_stats(self, xlsx_path):
        """从输出 xlsx 的“统计汇总” sheet 读取统计项"""
        if not os.path.isfile(xlsx_path):
            return
        try:
            import openpyxl
            wb = openpyxl.load_workbook(xlsx_path, read_only=True)
            if "统计汇总" not in wb.sheetnames:
                wb.close()
                return
            ws = wb["统计汇总"]
            stats = {}
            for row in ws.iter_rows(min_row=2, max_col=2, values_only=True):
                if row[0] is not None:
                    stats[str(row[0])] = row[1]
            wb.close()
        except Exception as e:
            self._append_log("[警告] 统计汇总读取失败: %s" % e)
            return
        self.stat_total.set_value(stats.get("总网格数", "--"))
        self.stat_precise.set_value(stats.get("精确区网格", "--"))
        self.stat_filled.set_value(stats.get("估计区网格(插值)", "--"))
        self.stat_volume.set_value(stats.get("总填方量(m³)", "--"))
